from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, extract
from app.api.dependencies import require_roles, get_tenant_db, get_current_user
from app.models.user import User
from app.models.finance import ExpenseClaim, TimeLog
from app.services.pdf_service import generate_executive_pdf
import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, date, timedelta

router = APIRouter()


@router.get("/export/executive")
async def export_executive_report(
    db: AsyncSession = Depends(get_tenant_db),
    current_user: dict = Depends(require_roles(["hr_admin", "sys_admin"]))
):
    """Agrupa datos de Empleados del Tenant y los formatea como PDF."""
    try:
        total_result = await db.execute(select(func.count()).select_from(User))
        total_users = total_result.scalar() or 0

        active_result = await db.execute(select(func.count()).select_from(User).where(User.is_active == True))
        active_users = active_result.scalar() or 0

        roles_result = await db.execute(
            select(User.role, func.count(User.id)).group_by(User.role)
        )
        roles_summary = []
        for role_name, count in roles_result.all():
            roles_summary.append({
                "role": (role_name or "Sin rol").replace("_", " ").title(),
                "count": count,
                "percentage": round((count / total_users * 100), 1) if total_users > 0 else 0
            })

        app_metadata = current_user.get("https://successcore.com/app_metadata", {})
        tenant_id = app_metadata.get("tenant_id") or current_user.get("tenant_id", "Unknown Co.")

        stats = {"total": total_users, "active": active_users, "roles": roles_summary}
        pdf_bytes = generate_executive_pdf(tenant_name=f"Data for {tenant_id}", stats=stats)

        headers = {'Content-Disposition': 'attachment; filename="executive_report.pdf"'}
        return Response(content=bytes(pdf_bytes), headers=headers, media_type="application/pdf")
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response(content=f"Error: {str(e)}", status_code=500, media_type="text/plain")


@router.get("/export/excel")
async def export_employees_excel(
    db: AsyncSession = Depends(get_tenant_db),
    current_user: dict = Depends(require_roles(["hr_admin", "sys_admin"]))
):
    """Exporta la lista completa de empleados del tenant como archivo .xlsx (US 4.1)."""
    result = await db.execute(select(User).order_by(User.created_at.desc()))
    employees = result.scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Empleados"

    HEADER_FILL = PatternFill("solid", fgColor="1E293B")
    HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    ALT_FILL = PatternFill("solid", fgColor="F1F5F9")
    BORDER = Border(
        bottom=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0")
    )

    col_headers = ["ID", "Nombre Completo", "Email", "Departamento", "Rol", "Estado", "Fecha Alta"]
    for col_num, header in enumerate(col_headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    ws.row_dimensions[1].height = 28

    for row_num, emp in enumerate(employees, 2):
        fill = ALT_FILL if row_num % 2 == 0 else PatternFill()
        row_data = [
            emp.id,
            emp.full_name or "—",
            emp.email,
            emp.department or "—",
            (emp.role or "Sin rol").replace("_", " ").title(),
            "Activo" if emp.is_active else "Archivado",
            emp.created_at.strftime("%d/%m/%Y") if emp.created_at else "—"
        ]
        for col_num, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center")

    col_widths = [36, 26, 34, 22, 18, 12, 14]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    excel_bytes = buffer.read()

    filename = f"empleados_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    resp_headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return Response(
        content=excel_bytes,
        headers=resp_headers,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/dashboard/data")
async def get_dashboard_data(
    metrics: str = Query("headcount", description="Comma-separated: headcount,departments,turnover,attendance,cost"),
    db: AsyncSession = Depends(get_tenant_db),
    _: dict = Depends(require_roles(["hr_admin", "sys_admin"]))
):
    metrics_list = [m.strip() for m in metrics.split(",")]
    data = {}

    total_result = await db.execute(select(func.count()).select_from(User))
    total = total_result.scalar() or 0

    if "headcount" in metrics_list:
        by_dept = await db.execute(
            select(User.department, func.count(User.id))
            .where(User.department != None)
            .group_by(User.department)
            .order_by(func.count(User.id).desc())
        )
        by_role = await db.execute(
            select(User.role, func.count(User.id)).group_by(User.role)
        )
        active_result = await db.execute(select(func.count()).select_from(User).where(User.is_active == True))
        data["headcount"] = {
            "total": total,
            "active": active_result.scalar() or 0,
            "inactive": total - (active_result.scalar() or 0),
            "by_department": [{"name": d or "Sin dept.", "count": c} for d, c in by_dept.all()],
            "by_role": [{"name": (r or "Sin rol").replace("_", " ").title(), "count": c} for r, c in by_role.all()],
        }

    if "departments" in metrics_list:
        dept_result = await db.execute(
            select(User.department, func.count(User.id))
            .where(User.department != None)
            .group_by(User.department)
            .order_by(func.count(User.id).desc())
        )
        data["departments"] = [
            {"name": d or "Sin dept.", "count": c, "percentage": round(c / total * 100, 1) if total > 0 else 0}
            for d, c in dept_result.all()
        ]

    if "turnover" in metrics_list:
        active_count = total
        inactive_count = 0
        for u in (await db.execute(select(User.is_active))).scalars().all():
            if not u:
                inactive_count += 1
        active_count = total - inactive_count
        rate = round(inactive_count / total * 100, 1) if total > 0 else 0.0
        data["turnover"] = {
            "rate": rate,
            "active": active_count,
            "inactive": inactive_count,
        }

    if "attendance" in metrics_list:
        today_start = datetime.combine(date.today(), datetime.min.time())
        today_count_result = await db.execute(
            select(func.count(func.distinct(TimeLog.user_id)))
            .where(TimeLog.clock_in >= today_start)
        )
        today_clocked = today_count_result.scalar() or 0
        today_rate = round(today_clocked / total * 100, 1) if total > 0 else 0.0
        trend = []
        for days_ago in range(6, -1, -1):
            day = date.today() - timedelta(days=days_ago)
            day_start = datetime.combine(day, datetime.min.time())
            day_end = datetime.combine(day + timedelta(days=1), datetime.min.time())
            count_result = await db.execute(
                select(func.count(func.distinct(TimeLog.user_id)))
                .where(TimeLog.clock_in >= day_start, TimeLog.clock_in < day_end)
            )
            count = count_result.scalar() or 0
            trend.append({
                "date": day.isoformat(),
                "count": count,
                "percentage": round(count / total * 100, 1) if total > 0 else 0.0,
            })
        data["attendance"] = {"rate": today_rate, "today_count": today_clocked, "total": total, "by_day": trend}

    if "cost" in metrics_list:
        month_start = date.today().replace(day=1)
        cost_result = await db.execute(
            select(func.coalesce(func.sum(ExpenseClaim.total_amount), 0))
            .where(ExpenseClaim.date >= month_start, ExpenseClaim.status != "rejected")
        )
        monthly_total = round(float(cost_result.scalar() or 0), 2)
        by_cat_result = await db.execute(
            select(ExpenseClaim.category, func.coalesce(func.sum(ExpenseClaim.total_amount), 0))
            .where(ExpenseClaim.date >= month_start, ExpenseClaim.status != "rejected")
            .group_by(ExpenseClaim.category)
            .order_by(func.sum(ExpenseClaim.total_amount).desc())
        )
        by_category = [
            {"category": c or "other", "amount": round(float(amt), 2)}
            for c, amt in by_cat_result.all()
        ]
        data["cost"] = {"total": monthly_total, "by_category": by_category, "currency": "EUR"}

    return data


@router.get("/export/custom")
async def export_custom_report(
    format: str = Query("csv", pattern="^(csv|pdf|excel)$"),
    db: AsyncSession = Depends(get_tenant_db),
    current_user: dict = Depends(require_roles(["hr_admin", "sys_admin"]))
):
    result = await db.execute(
        select(
            User.id, User.full_name, User.email, User.department, User.role,
            User.is_active, User.created_at
        ).order_by(User.created_at.desc())
    )
    rows = result.all()

    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Full Name", "Email", "Department", "Role", "Active", "Created At"])
        for row in rows:
            writer.writerow([
                row[0], row[1] or "", row[2], row[3] or "",
                (row[4] or "").replace("_", " ").title(),
                "Yes" if row[5] else "No",
                row[6].isoformat() if row[6] else ""
            ])
        content = output.getvalue().encode("utf-8-sig")
        filename = f"custom_report_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
        return Response(content=content, headers={"Content-Disposition": f'attachment; filename="{filename}"'}, media_type="text/csv")

    if format == "pdf":
        total = len(rows)
        active = sum(1 for r in rows if r[5])
        role_counts = {}
        for r in rows:
            role_name = (r[4] or "Sin rol").replace("_", " ").title()
            role_counts[role_name] = role_counts.get(role_name, 0) + 1
        roles_summary = [{"role": k, "count": v, "percentage": round(v / total * 100, 1) if total > 0 else 0} for k, v in role_counts.items()]
        stats = {"total": total, "active": active, "roles": roles_summary}
        app_metadata = current_user.get("https://successcore.com/app_metadata", {})
        tenant_id = app_metadata.get("tenant_id") or current_user.get("tenant_id", "Unknown Co.")
        pdf_bytes = generate_executive_pdf(tenant_name=f"Custom Report — {tenant_id}", stats=stats)
        filename = f"custom_report_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
        return Response(content=bytes(pdf_bytes), headers={"Content-Disposition": f'attachment; filename="{filename}"'}, media_type="application/pdf")

    if format == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Custom Report"
        HEADER_FILL = PatternFill("solid", fgColor="1E293B")
        HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
        ALT_FILL = PatternFill("solid", fgColor="F1F5F9")
        BORDER = Border(bottom=Side(style="thin", color="E2E8F0"), right=Side(style="thin", color="E2E8F0"))
        col_headers = ["ID", "Full Name", "Email", "Department", "Role", "Active", "Created At"]
        for col_num, header in enumerate(col_headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
        ws.row_dimensions[1].height = 28
        for row_num, row in enumerate(rows, 2):
            fill = ALT_FILL if row_num % 2 == 0 else PatternFill()
            row_data = [
                row[0], row[1] or "", row[2], row[3] or "",
                (row[4] or "").replace("_", " ").title(),
                "Yes" if row[5] else "No",
                row[6].strftime("%d/%m/%Y") if row[6] else ""
            ]
            for col_num, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_num, column=col_num, value=value)
                cell.fill = fill
                cell.border = BORDER
                cell.alignment = Alignment(vertical="center")
        col_widths = [36, 26, 34, 22, 18, 12, 14]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        filename = f"custom_report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        return Response(content=buffer.read(), headers={"Content-Disposition": f'attachment; filename="{filename}"'}, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    return Response(content="Unsupported format", status_code=400)


@router.get("/summary")
async def get_report_summary(
    db: AsyncSession = Depends(get_tenant_db),
    _: dict = Depends(require_roles(["hr_admin", "sys_admin"]))
):
    """Devuelve las métricas resumen para el dashboard de reportes (US 4.0)."""
    total_result = await db.execute(select(func.count()).select_from(User))
    total = total_result.scalar() or 0

    active_result = await db.execute(select(func.count()).select_from(User).where(User.is_active == True))
    active = active_result.scalar() or 0

    roles_result = await db.execute(
        select(User.role, func.count(User.id)).group_by(User.role)
    )
    roles = [
        {"name": r.title(), "count": c, "percentage": round(c / total * 100, 1) if total > 0 else 0}
        for r, c in roles_result.all()
    ]

    dept_result = await db.execute(
        select(User.department, func.count(User.id))
        .where(User.department != None)
        .group_by(User.department)
        .order_by(func.count(User.id).desc())
    )
    departments = [{"name": d or "Sin dept.", "count": c} for d, c in dept_result.all()]

    return {
        "total_employees": total,
        "active_employees": active,
        "inactive_employees": total - active,
        "roles_breakdown": roles,
        "departments": departments
    }


@router.get("/dashboard")
async def get_dashboard_summary(
    db: AsyncSession = Depends(get_tenant_db),
    _: dict = Depends(require_roles(["hr_admin", "sys_admin"]))
):
    """Devuelve todos los datos para el dashboard principal en una sola llamada."""
    total_result = await db.execute(select(func.count()).select_from(User))
    total = total_result.scalar() or 0

    active_result = await db.execute(select(func.count()).select_from(User).where(User.is_active == True))
    active = active_result.scalar() or 0

    dept_result = await db.execute(
        select(User.department, func.count(User.id))
        .where(User.department != None)
        .group_by(User.department)
        .order_by(func.count(User.id).desc())
        .limit(8)
    )
    departments = [{"name": d or "Sin dept.", "count": c} for d, c in dept_result.all()]

    roles_result = await db.execute(
        select(User.role, func.count(User.id)).group_by(User.role)
    )
    roles = [
        {"name": (r or "Sin rol").replace("_", " ").title(), "count": c, "percentage": round(c / total * 100, 1) if total > 0 else 0}
        for r, c in roles_result.all()
    ]

    recent_result = await db.execute(
        select(User).order_by(User.created_at.desc()).limit(5)
    )
    recent_employees = [
        {
            "id": u.id,
            "full_name": u.full_name or u.email,
            "department": u.department or "Sin departamento",
            "role": (u.role or "Sin rol").replace("_", " ").title(),
            "created_at": u.created_at.isoformat() if u.created_at else None
        }
        for u in recent_result.scalars().all()
    ]

    return {
        "total_employees": total,
        "active_employees": active,
        "inactive_employees": total - active,
        "unique_departments": len(departments),
        "departments": departments,
        "roles_breakdown": roles,
        "recent_employees": recent_employees,
    }
